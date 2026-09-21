import os
import datetime
import io
import re
import base64
import json
import hashlib
import urllib.request
import urllib.error
import pandas as pd
import numpy as np
import streamlit as st
import streamlit.components.v1 as components
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==========================================
# 1. ページ基本設定 & 明るい黄緑・薄緑ベースデザイン
# ==========================================
st.set_page_config(
    page_title="水処理点検・分析統合システム (KBL Wastewater Management)",
    page_icon="🌱",
    layout="wide",
    initial_sidebar_state="expanded"
)

# カスタムCSSスタイリング
st.markdown("""
<style>
    .main-header {
        font-size: 26px;
        font-weight: bold;
        color: #2E7D32;
        padding-bottom: 8px;
        border-bottom: 3.5px solid #66BB6A;
        margin-bottom: 15px;
    }
    .sub-header {
        font-size: 18px;
        font-weight: bold;
        color: #388E3C;
        margin-top: 15px;
        margin-bottom: 10px;
        padding-left: 8px;
        border-left: 4px solid #8BC34A;
    }
    .metric-card {
        background-color: #F1F8E9;
        border-radius: 8px;
        padding: 12px;
        border-left: 5px solid #66BB6A;
    }
    .stButton>button {
        background-color: #43A047;
        color: white;
        font-weight: bold;
        border-radius: 6px;
        border: none;
        padding: 8px 16px;
    }
    .stButton>button:hover {
        background-color: #66BB6A;
        color: white;
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 6px;
    }
    .stTabs [data-baseweb="tab"] {
        background-color: #F1F8E9;
        border-radius: 6px 6px 0px 0px;
        padding: 8px 18px;
        color: #388E3C;
        font-weight: bold;
    }
    .stTabs [aria-selected="true"] {
        background-color: #43A047 !important;
        color: white !important;
    }
    [data-testid="stSidebar"] {
        background-color: #F7FAF7;
    }
    .stAlert {
        border-radius: 8px;
    }
</style>
""", unsafe_allow_html=True)

# ------------------------------------------
# 安全な Secret 取得関数
# ------------------------------------------
def safe_get_secret(key, default=""):
    try:
        if hasattr(st, "secrets") and key in st.secrets:
            return st.secrets[key]
    except Exception:
        pass
    return default

correct_pwd = safe_get_secret("APP_PASSWORD", "kbl2026")

# ------------------------------------------
# 1.5 簡易パスワード認証保護 (20分以上セッション保持 & 復元機能)
# ------------------------------------------
if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False

now_ts = int(datetime.datetime.now().timestamp())
qp_time = st.query_params.get("auth_time", None)
qp_hash = st.query_params.get("auth_hash", None)

if not st.session_state["authenticated"] and qp_time and qp_hash:
    try:
        auth_ts = int(qp_time)
        expected_hash = hashlib.sha256(f"{correct_pwd}_{auth_ts}".encode()).hexdigest()[:16]
        if qp_hash == expected_hash and (now_ts - auth_ts) < 7200:
            st.session_state["authenticated"] = True
            new_hash = hashlib.sha256(f"{correct_pwd}_{now_ts}".encode()).hexdigest()[:16]
            st.query_params["auth_time"] = str(now_ts)
            st.query_params["auth_hash"] = new_hash
    except Exception:
        pass

if not st.session_state["authenticated"]:
    st.markdown("<div class='main-header'>🔒 KBL 排水管理システム - ログイン</div>", unsafe_allow_html=True)
    st.markdown("#### 🔑 パスワード認証")
    pwd_input = st.text_input("パスワード", type="password", key="login_pwd_key")
    if pwd_input:
        if pwd_input == correct_pwd:
            st.session_state["authenticated"] = True
            new_ts = int(datetime.datetime.now().timestamp())
            new_hash = hashlib.sha256(f"{correct_pwd}_{new_ts}".encode()).hexdigest()[:16]
            st.query_params["auth_time"] = str(new_ts)
            st.query_params["auth_hash"] = new_hash
            st.rerun()
        else:
            st.error("❌ パスワードが正しくありません。")
    st.stop()
else:
    if qp_time:
        try:
            if now_ts - int(qp_time) > 300:
                new_hash = hashlib.sha256(f"{correct_pwd}_{now_ts}".encode()).hexdigest()[:16]
                st.query_params["auth_time"] = str(now_ts)
                st.query_params["auth_hash"] = new_hash
        except Exception:
            pass

# バックグラウンド Ping ＆ スマホテンキー固定 (st.form の完全外側に配置)
components.html("""
<script>
setInterval(function() {
    try {
        fetch(window.location.href, { method: 'HEAD', mode: 'no-cors' }).catch(function(e){});
    } catch(e) {}
}, 30000);

function enforceNumericTenkey() {
    try {
        const inputs = window.parent.document.querySelectorAll('input[type="text"]');
        inputs.forEach(input => {
            input.setAttribute('inputmode', 'decimal');
            input.setAttribute('pattern', '[0-9.*#-]*');
        });
    } catch(e) {}
}
setTimeout(enforceNumericTenkey, 300);
setTimeout(enforceNumericTenkey, 1000);
setInterval(enforceNumericTenkey, 2000);
</script>
""", height=0)

# ------------------------------------------
# GitHub REST API 自動同期関数
# ------------------------------------------
def sync_to_github_api(file_path):
    try:
        github_token = safe_get_secret("GITHUB_TOKEN", "")
        if not github_token:
            return False, "st.secrets に GITHUB_TOKEN が設定されていません"
        
        repo = safe_get_secret("GITHUB_REPO", "kbl-wastewater-app")
        branch = safe_get_secret("GITHUB_BRANCH", "main")
        target_filename = os.path.basename(file_path)
        
        url = f"https://api.github.com/repos/{repo}/contents/{target_filename}"
        headers = {
            "Authorization": f"Bearer {github_token}",
            "Accept": "application/vnd.github.v3+json",
            "User-Agent": "StreamlitApp"
        }
        
        sha = None
        req_get = urllib.request.Request(f"{url}?ref={branch}", headers=headers, method="GET")
        try:
            with urllib.request.urlopen(req_get) as resp:
                res_data = json.loads(resp.read().decode("utf-8"))
                sha = res_data.get("sha")
        except urllib.error.HTTPError as e:
            if e.code != 404:
                return False, f"HTTP Error {e.code} (SHA取得失敗)"
                
        with open(file_path, "rb") as f:
            content_b64 = base64.b64encode(f.read()).decode("utf-8")
            
        payload = {
            "message": f"Auto-update wastewater data via App [{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}]",
            "content": content_b64,
            "branch": branch
        }
        if sha:
            payload["sha"] = sha
            
        data_json = json.dumps(payload).encode("utf-8")
        req_put = urllib.request.Request(url, data=data_json, headers=headers, method="PUT")
        
        with urllib.request.urlopen(req_put) as resp:
            if resp.status in [200, 201]:
                return True, "成功"
            return False, f"HTTP Status {resp.status}"
    except Exception as ex:
        return False, str(ex)

# ------------------------------------------
# 自動計算関数 (有機割合、無機割合、沈降速度、水面積負荷、返送率、SVI、BOD除去率)
# ------------------------------------------
def calculate_auto_metrics(df_item_all, df_loc_all, selected_site_id, selected_sheet_type, form_data_dict, df_rec_latest, target_rep_id):
    calc_updates = {}
    
    def _parse_num(item_id):
        if item_id in form_data_dict and form_data_dict[item_id] is not None:
            v_str = str(form_data_dict[item_id]).replace(',', '').replace('%', '').replace('<', '').replace('>', '').replace('未満', '').replace('以上', '').strip()
            try:
                return float(v_str)
            except ValueError:
                pass
        match = df_rec_latest[(df_rec_latest['Report_ID'] == target_rep_id) & (df_rec_latest['Item_ID'] == item_id)]
        if not match.empty:
            v_str = str(match.iloc[0]['Value']).replace(',', '').replace('%', '').replace('<', '').replace('>', '').replace('未満', '').replace('以上', '').strip()
            try:
                return float(v_str)
            except ValueError:
                pass
        return None

    def _get_raw_val(item_id):
        if item_id in form_data_dict and form_data_dict[item_id] is not None:
            return str(form_data_dict[item_id]).strip()
        match = df_rec_latest[(df_rec_latest['Report_ID'] == target_rep_id) & (df_rec_latest['Item_ID'] == item_id)]
        if not match.empty:
            return str(match.iloc[0]['Value']).strip()
        return ""

    def _find_items(df_subset, query_str):
        exact = df_subset[df_subset['Item_Name'] == query_str]
        if not exact.empty:
            return exact
        prefix = df_subset[df_subset['Item_Name'].astype(str).str.startswith(query_str)]
        if not prefix.empty:
            return prefix
        contains = df_subset[df_subset['Item_Name'].astype(str).str.contains(query_str, regex=False)]
        return contains

    # 1. 有機割合 [%] (MLVSS/MLSS*100) & 無機割合 [%] (100 - 有機割合)
    locs_in_site = df_loc_all[df_loc_all['Site_ID'] == selected_site_id]['Loc_ID'].unique()
    for loc_id in locs_in_site:
        loc_items = df_item_all[(df_item_all['Loc_ID'] == loc_id) & (df_item_all['Sheet_Type'] == selected_sheet_type)]
        mlss_item = _find_items(loc_items, 'MLSS')
        if mlss_item.empty:
            mlss_item = _find_items(loc_items, 'MLSS(簡)')
        mlvss_item = _find_items(loc_items, 'MLVSS')
        org_item = _find_items(loc_items, '有機割合')
        inorg_item = _find_items(loc_items, '無機割合')
        
        if not mlss_item.empty and not mlvss_item.empty:
            mlss_v = _parse_num(mlss_item.iloc[0]['Item_ID'])
            mlvss_v = _parse_num(mlvss_item.iloc[0]['Item_ID'])
            if mlss_v and mlss_v > 0 and mlvss_v is not None:
                org_ratio = round((mlvss_v / mlss_v) * 100.0, 1)
                inorg_ratio = round(100.0 - org_ratio, 1)
                if not org_item.empty:
                    calc_updates[org_item.iloc[0]['Item_ID']] = f"{org_ratio}%"
                if not inorg_item.empty:
                    calc_updates[inorg_item.iloc[0]['Item_ID']] = f"{inorg_ratio}%"

    # 2. SVI Calculation (点検管理表: SV30 * 10000 / MLSS)
    sv30_mlss_map = {
        'S001': ('I0028', 'I0027', 'I0231'),
        'S002': ('I0075', 'I0074', 'I0232'),
        'S004': ('I0137', 'I0136', 'I0233'),
        'S005': ('I0204', 'I0203', 'I0234')
    }
    if selected_sheet_type == '点検管理表':
        if selected_site_id in sv30_mlss_map:
            sv30_id, mlss_id, svi_id = sv30_mlss_map[selected_site_id]
            sv30_v = _parse_num(sv30_id)
            mlss_v = _parse_num(mlss_id)
            if sv30_v is not None and sv30_v > 0 and mlss_v is not None and mlss_v > 0:
                svi_val = round((sv30_v * 10000.0) / mlss_v, 1)
                calc_updates[svi_id] = str(svi_val)
        else:
            # Dynamic fallback
            site_items = df_item_all.merge(df_loc_all[df_loc_all['Site_ID'] == selected_site_id], on='Loc_ID')
            site_items = site_items[site_items['Sheet_Type'] == '点検管理表']
            sv30_item = _find_items(site_items, 'SV30')
            mlss_item = _find_items(site_items, 'MLSS')
            if mlss_item.empty:
                mlss_item = _find_items(site_items, 'MLSS(簡)')
            svi_item = _find_items(site_items, 'SVI')
            if not sv30_item.empty and not mlss_item.empty and not svi_item.empty:
                sv30_v = _parse_num(sv30_item.iloc[0]['Item_ID'])
                mlss_v = _parse_num(mlss_item.iloc[0]['Item_ID'])
                if sv30_v is not None and sv30_v > 0 and mlss_v is not None and mlss_v > 0:
                    svi_val = round((sv30_v * 10000.0) / mlss_v, 1)
                    calc_updates[svi_item.iloc[0]['Item_ID']] = str(svi_val)

    # 3. BOD除去率 Calculation: (原水BOD - 処理水BOD) / 原水BOD * 100
    bod_rem_map = {
        'S001': [
            ('I0049', 'I0054', 'I0235'),  # 計量証明
            ('I0003', 'I0045', 'I0238')   # 点検管理表
        ],
        'S002': [
            ('I0090', 'I0108', 'I0236')   # 計量証明
        ],
        'S004': [
            ('I0160', 'I0177', 'I0237')   # 計量証明
        ]
    }
    if selected_site_id in bod_rem_map:
        for raw_bod_id, treated_bod_id, rem_item_id in bod_rem_map[selected_site_id]:
            target_item_row = df_item_all[df_item_all['Item_ID'] == rem_item_id]
            if not target_item_row.empty and target_item_row.iloc[0]['Sheet_Type'] == selected_sheet_type:
                raw_v = _parse_num(raw_bod_id)
                treated_raw_str = _get_raw_val(treated_bod_id)
                if raw_v is not None and raw_v > 0 and treated_raw_str and treated_raw_str not in ['', '-', 'nan']:
                    is_less = ('未満' in treated_raw_str) or ('<' in treated_raw_str)
                    treated_clean = treated_raw_str.replace(',', '').replace('%', '').replace('<', '').replace('>', '').replace('未満', '').replace('以上', '').strip()
                    treated_v = None
                    try:
                        treated_v = float(treated_clean)
                    except ValueError:
                        pass
                        
                    if treated_v is not None:
                        if is_less or treated_v < 5.0:
                            rem_rate = round((raw_v - 5.0) / raw_v * 100.0, 1)
                            calc_updates[rem_item_id] = f"{rem_rate}%以上"
                        else:
                            rem_rate = round((raw_v - treated_v) / raw_v * 100.0, 1)
                            calc_updates[rem_item_id] = f"{rem_rate}%"

    # 4. キャンパック（S004/S005）専用計算 (沈降速度, 水面積負荷, 沈降速度/水面積負荷, 返送率)
    if selected_site_id in ['S004', 'S005'] and selected_sheet_type == '点検管理表':
        site_items = df_item_all.merge(df_loc_all[df_loc_all['Site_ID'] == selected_site_id], on='Loc_ID')
        site_items = site_items[site_items['Sheet_Type'] == '点検管理表']
        
        isou_item = _find_items(site_items, '移送量')
        hensou_item = _find_items(site_items, '返送量')
        
        aeration_end = site_items[site_items['Loc_Name'] == '曝気槽（末端側）']
        mlss_item = _find_items(aeration_end, 'MLSS')
        if mlss_item.empty:
            mlss_item = _find_items(aeration_end, 'MLSS(簡)')
        temp_item = _find_items(aeration_end, '水温')
        sv30_item = _find_items(aeration_end, 'SV30')
        
        isou_v = _parse_num(isou_item.iloc[0]['Item_ID']) if not isou_item.empty else None
        hensou_v = _parse_num(hensou_item.iloc[0]['Item_ID']) if not hensou_item.empty else None
        mlss_v = _parse_num(mlss_item.iloc[0]['Item_ID']) if not mlss_item.empty else None
        temp_v = _parse_num(temp_item.iloc[0]['Item_ID']) if not temp_item.empty else None
        sv30_v = _parse_num(sv30_item.iloc[0]['Item_ID']) if not sv30_item.empty else None
        
        load_item = _find_items(site_items, '水面積負荷')
        return_ratio_item = _find_items(site_items, '返送率')
        settling_item = _find_items(site_items, '沈降速度')
        ratio_item = site_items[site_items['Item_Name'] == '沈降速度/水面積負荷']
        if ratio_item.empty:
            ratio_item = _find_items(site_items, '沈降速度/水面積負荷')
            
        if not settling_item.empty:
            settling_item = settling_item[~settling_item['Item_Name'].astype(str).str.contains('水面積負荷')]
        
        surf_load = None
        if isou_v is not None and isou_v > 0:
            surf_load = round(isou_v / 143.0, 2)
            if not load_item.empty:
                calc_updates[load_item.iloc[0]['Item_ID']] = str(surf_load)
                
        if isou_v is not None and isou_v > 0 and hensou_v is not None and hensou_v >= 0:
            r_ratio = round((hensou_v / isou_v) * 100.0, 1)
            if not return_ratio_item.empty:
                calc_updates[return_ratio_item.iloc[0]['Item_ID']] = f"{r_ratio}%"
                
        v_settling = None
        if mlss_v is not None and mlss_v > 0 and temp_v is not None and temp_v > 0 and sv30_v is not None and sv30_v > 0:
            try:
                sv_term = (sv30_v * 10000.0) / mlss_v
                v_settling = (1.78 * (10**7) * (mlss_v**(-1.46)) * (temp_v**(0.853)) * (sv_term**(-0.804))) / 24.0
                v_settling = round(v_settling, 2)
                if not settling_item.empty:
                    calc_updates[settling_item.iloc[0]['Item_ID']] = str(v_settling)
            except Exception:
                pass
                
        if v_settling is not None and surf_load is not None and surf_load > 0:
            ratio_val = round(v_settling / surf_load, 2)
            if not ratio_item.empty:
                calc_updates[ratio_item.iloc[0]['Item_ID']] = str(ratio_val)
                
    return calc_updates

# ------------------------------------------
# データベース探射・自動作成・堅牢ロード処理
# ------------------------------------------
def find_excel_db():
    for fname in ["wastewater-appsheet-db-v3.xlsx", "wastewater-appsheet-db-v2.xlsx", "wastewater-appsheet-db.xlsx"]:
        if os.path.exists(fname):
            return fname
        abs_p = f"/workspace/artifacts/{fname}"
        if os.path.exists(abs_p):
            return abs_p
            
    try:
        for root, _, files in os.walk("."):
            for file in files:
                if file.endswith(".xlsx") and "wastewater" in file:
                    return os.path.join(root, file)
    except Exception:
        pass
                
    try:
        if os.path.exists("/workspace/artifacts"):
            for root, _, files in os.walk("/workspace/artifacts"):
                for file in files:
                    if file.endswith(".xlsx") and "wastewater" in file:
                        return os.path.join(root, file)
    except Exception:
        pass
                    
    return None

db_path = find_excel_db()

@st.cache_data(ttl=1)
def load_all_data(path):
    if not path or not os.path.exists(path):
        return None, None, None, None, None
        
    try:
        xls = pd.ExcelFile(path)
        df_site = pd.read_excel(xls, "Site Master")
        df_loc = pd.read_excel(xls, "Location Master")
        df_item = pd.read_excel(xls, "Item Master")
        df_rep = pd.read_excel(xls, "Daily Report")
        df_rec = pd.read_excel(xls, "Inspection Records")
        
        df_site = df_site[~df_site["Site_Name"].astype(str).str.contains("北越") & (df_site["Site_ID"] != "S003")]
        df_loc = df_loc[~df_loc["Site_ID"].isin(["S003"]) & df_loc["Site_ID"].isin(df_site["Site_ID"])]
        df_item = df_item[df_item["Loc_ID"].isin(df_loc["Loc_ID"])]
        df_rep = df_rep[df_rep["Site_ID"].isin(df_site["Site_ID"])]
        df_rec = df_rec[df_rec["Report_ID"].isin(df_rep["Report_ID"])]

        df_rep["Date"] = pd.to_datetime(df_rep["Date"]).dt.strftime("%Y/%m/%d")
        
        val_str_series = df_rec["Value"].astype(str)
        val_clean = (
            val_str_series
            .str.replace(",", "", regex=False)
            .str.replace("%", "", regex=False)
            .str.replace("<", "", regex=False)
            .str.replace(">", "", regex=False)
            .str.replace("未満", "", regex=False)
            .str.replace("以上", "", regex=False)
            .str.strip()
        )
        df_rec["Value_Num"] = pd.to_numeric(val_clean, errors="coerce")
        return df_site, df_loc, df_item, df_rep, df_rec
    except Exception as ex:
        st.error(f"Excelデータベースの読み込み中にエラーが発生しました: {ex}")
        return None, None, None, None, None

loaded_data = load_all_data(db_path) if db_path else (None, None, None, None, None)
df_site, df_loc, df_item, df_rep, df_rec = loaded_data

if df_site is None:
    st.markdown("<div class='main-header'>🌱 KBL 排水処理統合管理システム</div>", unsafe_allow_html=True)
    st.warning("⚠️ **データベースファイル (`wastewater-appsheet-db-v3.xlsx`) が自動検出されませんでした。**")
    st.info("以下からデータベースExcelファイルをアップロードするか、GitHubリポジトリに `wastewater-appsheet-db-v3.xlsx` をコミットしてください。")
    
    uploaded_db = st.file_uploader("📁 データベースExcelファイルをアップロード", type=["xlsx"])
    if uploaded_db is not None:
        target_save_p = "wastewater-appsheet-db-v3.xlsx"
        with open(target_save_p, "wb") as f:
            f.write(uploaded_db.getbuffer())
        st.success("✅ データベースのアップロードが完了しました！再読み込みします...")
        st.cache_data.clear()
        st.rerun()
    st.stop()

# サイドバー設定
st.sidebar.markdown("<h2 style='color:#2E7D32;'>⚙️ KBL 排水管理設定</h2>", unsafe_allow_html=True)

site_options = dict(zip(df_site["Site_ID"], df_site["Site_Name"]))
selected_site_id = st.sidebar.selectbox("① 対象現場（工場）を選択", options=list(site_options.keys()), format_func=lambda x: site_options[x])

sheet_type_options = ["点検管理表", "計量証明"]
selected_sheet_type = st.sidebar.selectbox("② 管理シート種別を選択", options=sheet_type_options)

st.sidebar.markdown("---")
site_name_disp = site_options[selected_site_id]
st.sidebar.info(f"📍 選択中現場: **{site_name_disp}**\n📄 種別: **{selected_sheet_type}**")

# メインタイトル
st.markdown(f"<div class='main-header'>🌱 KBL 排水処理統合管理システム - {site_options[selected_site_id]}</div>", unsafe_allow_html=True)

# ------------------------------------------
# 4. タブ構築
# ------------------------------------------
tab1, tab2, tab3, tab4 = st.tabs([
    "📝 点検データ入力・異常値チェック",
    "📊 水質比較グラフ",
    "📑 ⑥ エクセル帳票ダウンロード",
    "🗃️ データベース全件閲覧"
])

# ------------------------------------------
# TAB 1: 点検データ入力 + 異常値チェック
# ------------------------------------------
with tab1:
    st.markdown("<div class='sub-header'>📝 点検結果の新規入力 ＆ 蓄積 (過去データとの異常値チェック機能付き)</div>", unsafe_allow_html=True)
    
    if "save_success_msg" in st.session_state:
        st.success(st.session_state["save_success_msg"])
        del st.session_state["save_success_msg"]

    col_input1, col_input2 = st.columns(2)

    with col_input1:
        input_date = st.date_input("③ 点検日を選択", datetime.date.today())
        input_date_str = input_date.strftime("%Y/%m/%d")

    with col_input2:
        locs_with_items = df_item[(df_item["Loc_ID"].isin(df_loc[df_loc["Site_ID"] == selected_site_id]["Loc_ID"])) & (df_item["Sheet_Type"] == selected_sheet_type)]["Loc_ID"].unique()
        site_locs = df_loc[(df_loc["Site_ID"] == selected_site_id) & (df_loc["Loc_ID"].isin(locs_with_items))].sort_values("Display_Order")
        if site_locs.empty:
            st.warning(f"⚠️ 【{site_options[selected_site_id]}】には「{selected_sheet_type}」の項目データが登録されていません。")
            selected_loc_id = None
        else:
            loc_options = dict(zip(site_locs["Loc_ID"], site_locs["Loc_Name"]))
            selected_loc_id = st.selectbox("④ 槽（測定箇所）を選択", options=list(loc_options.keys()), format_func=lambda x: loc_options[x], key="input_loc")

    st.markdown("---")
    if selected_loc_id:
        st.markdown(f"#### 📍 【{site_options[selected_site_id]}】 - 『{loc_options[selected_loc_id]}』 ({selected_sheet_type})")

        loc_items = df_item[(df_item["Loc_ID"] == selected_loc_id) & (df_item["Sheet_Type"] == selected_sheet_type)].sort_values("Display_Order")

        if loc_items.empty:
            st.warning("指定された槽・シート種別の点検項目が登録されていません。")
        else:
            existing_rep = df_rep[(df_rep["Site_ID"] == selected_site_id) & (df_rep["Date"] == input_date_str) & (df_rep["Sheet_Type"] == selected_sheet_type)]
            existing_values = {}
            if not existing_rep.empty:
                rep_id_exist = existing_rep.iloc[0]["Report_ID"]
                rec_exist = df_rec[df_rec["Report_ID"] == rep_id_exist]
                existing_values = dict(zip(rec_exist["Item_ID"], rec_exist["Value"]))
                st.info(f"💡 {input_date_str} の既存データが読み込まれました。必要に応じて内容を更新してください。")

            df_hist = df_rec[df_rec["Item_ID"].isin(loc_items["Item_ID"])].dropna(subset=["Value_Num"])
            stats_by_item = {}
            for item_id_key, group in df_hist.groupby("Item_ID"):
                vals = group["Value_Num"].values
                if len(vals) >= 2:
                    mean_v = np.mean(vals)
                    std_v = np.std(vals)
                    min_v = np.min(vals)
                    max_v = np.max(vals)
                    lower_b = max(0, mean_v - 2.5 * std_v) if mean_v >= 0 else mean_v - 2.5 * std_v
                    upper_b = mean_v + 2.5 * std_v
                    stats_by_item[item_id_key] = {
                        "mean": mean_v, "std": std_v, "min": min_v, "max": max_v,
                        "lower": lower_b, "upper": upper_b
                    }
                elif len(vals) == 1:
                    stats_by_item[item_id_key] = {
                        "mean": vals[0], "std": 0, "min": vals[0], "max": vals[0],
                        "lower": vals[0]*0.5, "upper": vals[0]*1.5
                    }

            st.info("💡 **自動計算項目（SVI・BOD除去率・有機割合・無機割合・水面積負荷・沈降速度・返送率など）** は、入力された基礎数値から保存時に**全自動計算されExcelデータベースへ直接格納**されます（入力ボックスは非表示としています）。")
            
            calc_item_names = ['有機割合', '有機割合 (%)', '無機割合', '沈降速度', '水面積負荷', '沈降速度/水面積負荷', '返送率', '返送率 (%)', 'SVI', 'BOD除去率']
            input_items = loc_items[~loc_items['Item_Name'].isin(calc_item_names)].sort_values("Display_Order")

            with st.form("inspection_input_form"):
                form_data = {}
                input_items_list = list(input_items.iterrows())

                # スマホ1列折りたたみ時順序 (A1 -> B1 -> C1 -> A2 -> B2 -> C2) を保証する3項目区切り行構成
                for row_start in range(0, len(input_items_list), 3):
                    row_chunk = input_items_list[row_start : row_start + 3]
                    cols = st.columns(3)
                    for c_idx, (_, item) in enumerate(row_chunk):
                        item_id = item["Item_ID"]
                        item_name = item["Item_Name"]
                        vtype = item["Value_Type"]
                        unit = str(item["Unit_or_Options"]) if pd.notna(item["Unit_or_Options"]) else ""

                        default_val = existing_values.get(item_id, "")
                        label_str = item_name if vtype == "Enum" or not unit or unit == "nan" else f"{item_name} ({unit})"
                        col_target = cols[c_idx]

                        with col_target:
                            if vtype == "Enum":
                                opts = ["-", "無", "微少", "少", "中", "多"]
                                curr_idx = opts.index(default_val) if default_val in opts else 0
                                form_data[item_id] = st.selectbox(label_str, options=opts, index=curr_idx)
                            else:
                                init_str = str(default_val) if str(default_val) != "nan" else ""
                                form_data[item_id] = st.text_input(label_str, value=init_str, key=f"inp_{item_id}")
                                if item_id in stats_by_item:
                                    st_info = stats_by_item[item_id]
                                    st.caption(f"💡 過去平均: {st_info['mean']:.2f} (目安: {st_info['lower']:.1f} 〜 {st_info['upper']:.1f})")

                notes_input = st.text_area("備考 (Notes)", value=existing_rep.iloc[0]["Notes"] if not existing_rep.empty and pd.notna(existing_rep.iloc[0]["Notes"]) else "")

                submit_btn = st.form_submit_button("💾 点検データを保存・蓄積する")

                if submit_btn:
                    anomalies = []
                    for item_id, val in form_data.items():
                        if val is not None:
                            val_str = str(val).strip().replace(",", "").replace("%", "").replace("<", "").replace(">", "")
                            if val_str and val_str not in ["-", "nan", "None"]:
                                try:
                                    num_v = float(val_str)
                                    if item_id in stats_by_item:
                                        st_info = stats_by_item[item_id]
                                        if num_v < st_info['lower'] or num_v > st_info['upper']:
                                            it_name = df_item[df_item["Item_ID"] == item_id].iloc[0]["Item_Name"]
                                            anomalies.append(f"・**{it_name}**: 入力値 {num_v} (過去通常範囲: {st_info['lower']:.1f} 〜 {st_info['upper']:.1f} / 平均: {st_info['mean']:.2f})")
                                except ValueError:
                                    pass

                    if anomalies:
                        st.warning("⚠️ **【異常値・誤入力の可能性あり】**\n" + "\n".join(anomalies) + "\n\n※数値に問題が無ければ、データは正常に保存されます。")

                    wb = openpyxl.load_workbook(db_path)
                    ws_rep = wb["Daily Report"]
                    ws_rec = wb["Inspection Records"]

                    df_rep_latest = pd.read_excel(db_path, "Daily Report")
                    df_rec_latest = pd.read_excel(db_path, "Inspection Records")

                    exist_rep_match = df_rep_latest[(df_rep_latest["Site_ID"] == selected_site_id) & (df_rep_latest["Date"] == input_date_str) & (df_rep_latest["Sheet_Type"] == selected_sheet_type)]

                    if not exist_rep_match.empty:
                        target_rep_id = exist_rep_match.iloc[0]["Report_ID"]
                        for row in ws_rep.iter_rows(min_row=2):
                            if str(row[0].value) == str(target_rep_id):
                                row[4].value = ""
                                row[5].value = notes_input
                                break
                    else:
                        max_rep_num = 0
                        for rid in df_rep_latest["Report_ID"].dropna():
                            m = re.search(r"\d+", str(rid))
                            if m:
                                max_rep_num = max(max_rep_num, int(m.group()))
                        target_rep_id = f"REP{max_rep_num+1:05d}"
                        ws_rep.append([target_rep_id, input_date_str, selected_site_id, selected_sheet_type, "", notes_input])

                    existing_item_ids = set(df_rec_latest[df_rec_latest["Report_ID"] == target_rep_id]["Item_ID"])

                    max_rec_num = 0
                    for rcid in df_rec_latest["Rec_ID"].dropna():
                        m = re.search(r"\d+", str(rcid))
                        if m:
                            max_rec_num = max(max_rec_num, int(m.group()))

                    auto_calcs = calculate_auto_metrics(df_item, df_loc, selected_site_id, selected_sheet_type, form_data, df_rec_latest, target_rep_id)
                    final_save_data = form_data.copy()
                    final_save_data.update(auto_calcs)

                    for item_id, val in final_save_data.items():
                        if val is None or str(val).strip() in ["", "-", "nan", "None"]:
                            continue
                        if item_id in existing_item_ids:
                            for row in ws_rec.iter_rows(min_row=2):
                                if str(row[1].value) == str(target_rep_id) and str(row[2].value) == str(item_id):
                                    row[3].value = str(val).strip()
                                    break
                        else:
                            max_rec_num += 1
                            new_rec_id = f"REC{max_rec_num:06d}"
                            ws_rec.append([new_rec_id, target_rep_id, item_id, str(val).strip()])

                    wb.save(db_path)
                    st.cache_data.clear()
                    
                    sync_ok, sync_msg = sync_to_github_api(db_path)
                    if sync_ok:
                        st.session_state["save_success_msg"] = f"✅ {input_date_str} 『{loc_options[selected_loc_id]}』 の点検データを正常に保存しました（GitHubへの自動同期も成功しました）！"
                    else:
                        st.session_state["save_success_msg"] = f"✅ {input_date_str} 『{loc_options[selected_loc_id]}』 の点検データを正常に保存しました（ローカル更新完了 / GitHub同期: {sync_msg}）"
                    st.rerun()

# ------------------------------------------
# TAB 2: 動的2軸・複数槽比較水質グラフ
# ------------------------------------------
with tab2:
    st.markdown("<div class='sub-header'>📊 期間・複数槽(最大3箇所)・1軸/2軸の比較水質グラフ</div>", unsafe_allow_html=True)

    g_locs_with_items = df_item[(
        df_item["Loc_ID"].isin(df_loc[df_loc["Site_ID"] == selected_site_id]["Loc_ID"])
    ) & (df_item["Sheet_Type"] == selected_sheet_type)]["Loc_ID"].unique()

    g_site_locs = df_loc[(
        df_loc["Site_ID"] == selected_site_id
    ) & (df_loc["Loc_ID"].isin(g_locs_with_items))].sort_values("Display_Order")

    if g_site_locs.empty:
        st.warning("⚠️ 該当する槽・測定箇所データが存在しません。")
    else:
        graph_loc_options = dict(zip(g_site_locs["Loc_ID"], g_site_locs["Loc_Name"]))
        all_loc_ids = list(graph_loc_options.keys())

        col_g1, col_g2, col_g3 = st.columns([2.5, 2, 2])

        with col_g1:
            selected_graph_locs = st.multiselect(
                "📍 比較する槽（測定箇所）を選択 (最大3つ)",
                options=all_loc_ids,
                default=all_loc_ids[:min(3, len(all_loc_ids))],
                format_func=lambda x: graph_loc_options[x],
                max_selections=3,
                key="g_multilocs"
            )

        if not selected_graph_locs:
            st.info("比較する槽を1つ以上選択してください。")
        else:
            sel_items = df_item[(
                df_item["Loc_ID"].isin(selected_graph_locs)
            ) & (df_item["Sheet_Type"] == selected_sheet_type)].sort_values("Display_Order")

            unique_item_names = []
            item_name_to_unit = {}
            for _, r in sel_items.iterrows():
                iname = r["Item_Name"]
                uopt = str(r["Unit_or_Options"]) if pd.notna(r["Unit_or_Options"]) else ""
                if iname not in unique_item_names:
                    unique_item_names.append(iname)
                if iname in ["無機割合", "有機割合", "返送率", "BOD除去率"]:
                    item_name_to_unit[iname] = " (%)"
                elif uopt and uopt != "nan":
                    item_name_to_unit[iname] = f" ({uopt})"
                elif iname not in item_name_to_unit:
                    item_name_to_unit[iname] = ""

            item_display_map = {iname: f"{iname}{item_name_to_unit.get(iname, '')}" for iname in unique_item_names}

            with col_g2:
                selected_item_y1 = st.selectbox(
                    "📈 第1縦軸 (左軸/実線・●)",
                    options=unique_item_names,
                    format_func=lambda x: item_display_map[x],
                    key="g_y1_name"
                )

            with col_g3:
                y2_options = [None] + unique_item_names
                selected_item_y2 = st.selectbox(
                    "📉 第2縦軸 (右軸/実線・◆ - オプション)",
                    options=y2_options,
                    format_func=lambda x: "なし (1軸のみ)" if x is None else item_display_map[x],
                    key="g_y2_name"
                )

            st.markdown("---")
            st.markdown("##### 📅 グラフ表示期間の設定")

            site_reps = df_rep[(df_rep["Site_ID"] == selected_site_id) & (df_rep["Sheet_Type"] == selected_sheet_type)]
            if not site_reps.empty:
                site_dates_dt = pd.to_datetime(site_reps["Date"]).sort_values().drop_duplicates()
                all_date_strs = site_dates_dt.dt.strftime("%Y/%m/%d").tolist()
                min_d = site_dates_dt.min().date()
                max_d = site_dates_dt.max().date()
            else:
                all_date_strs = [datetime.date.today().strftime("%Y/%m/%d")]
                min_d = datetime.date(2023, 1, 1)
                max_d = datetime.date.today()

            preset_opt = st.radio(
                "⏱️ 期間プリセット選択",
                options=["全期間", "過去3ヶ月", "過去6ヶ月", "過去1年", "過去2年", "日付範囲を個別指定"],
                index=0,
                horizontal=True,
                key="period_preset_radio"
            )

            if preset_opt == "全期間":
                start_d = min_d
                end_d = max_d
            elif preset_opt == "過去3ヶ月":
                start_d = max(min_d, (pd.to_datetime(max_d) - pd.DateOffset(months=3)).date())
                end_d = max_d
            elif preset_opt == "過去6ヶ月":
                start_d = max(min_d, (pd.to_datetime(max_d) - pd.DateOffset(months=6)).date())
                end_d = max_d
            elif preset_opt == "過去1年":
                start_d = max(min_d, (pd.to_datetime(max_d) - pd.DateOffset(years=1)).date())
                end_d = max_d
            elif preset_opt == "過去2年":
                start_d = max(min_d, (pd.to_datetime(max_d) - pd.DateOffset(years=2)).date())
                end_d = max_d
            else:
                d_col1, d_col2 = st.columns(2)
                with d_col1:
                    start_d = st.date_input("📅 開始点検日", value=min_d, min_value=min_d, max_value=max_d, key="custom_start_d")
                with d_col2:
                    end_d = st.date_input("📅 終了点検日", value=max_d, min_value=min_d, max_value=max_d, key="custom_end_d")
                if isinstance(start_d, (list, tuple)):
                    start_d = start_d[0] if len(start_d) > 0 else min_d
                if isinstance(end_d, (list, tuple)):
                    end_d = end_d[-1] if len(end_d) > 0 else max_d

            active_dates_in_range = [d for d in site_dates_dt.dt.date if start_d <= d <= end_d]
            st.info(f"💡 **選択中の期間 ({start_d.strftime('%Y/%m/%d')} 〜 {end_d.strftime('%Y/%m/%d')}) に含まれる点検日**: **{len(active_dates_in_range)} 回**")

            st.markdown("---")

            site_reps_filtered = df_rep[(
                df_rep["Site_ID"] == selected_site_id
            ) & (
                df_rep["Sheet_Type"] == selected_sheet_type
            )].copy()
            site_reps_filtered["Date_dt"] = pd.to_datetime(site_reps_filtered["Date"])

            if start_d and end_d:
                site_reps_filtered = site_reps_filtered[
                    (site_reps_filtered["Date_dt"].dt.date >= start_d) &
                    (site_reps_filtered["Date_dt"].dt.date <= end_d)
                ]
            site_reps_filtered = site_reps_filtered.sort_values("Date_dt")

            loc_colors = [
                {"y1": "#2E7D32", "y2": "#1976D2", "label": "槽1(緑/青)"},
                {"y1": "#7CB342", "y2": "#8E24AA", "label": "槽2(黄緑/紫)"},
                {"y1": "#F57C00", "y2": "#00ACC1", "label": "槽3(橙/水色)"}
            ]

            fig = make_subplots(specs=[[{"secondary_y": True}]])
            has_data = False
            summary_data = []

            for l_idx, loc_id in enumerate(selected_graph_locs):
                loc_name = graph_loc_options[loc_id]
                color_cfg = loc_colors[l_idx % len(loc_colors)]

                item_y1_row = df_item[(
                    df_item["Loc_ID"] == loc_id
                ) & (
                    df_item["Sheet_Type"] == selected_sheet_type
                ) & (
                    df_item["Item_Name"] == selected_item_y1
                )]

                if not item_y1_row.empty:
                    item_y1_id = item_y1_row.iloc[0]["Item_ID"]
                    rec_y1 = df_rec[df_rec["Item_ID"] == item_y1_id][["Report_ID", "Value_Num"]]
                    df_plot_y1 = site_reps_filtered.merge(rec_y1, on="Report_ID", how="left")

                    valid_y1 = df_plot_y1.dropna(subset=["Value_Num"])
                    if not valid_y1.empty:
                        has_data = True
                        fig.add_trace(
                            go.Scatter(
                                x=df_plot_y1["Date"],
                                y=df_plot_y1["Value_Num"],
                                name=f"【{loc_name}】{item_display_map[selected_item_y1]}",
                                mode="lines+markers",
                                connectgaps=False,
                                line=dict(color=color_cfg["y1"], width=3),
                                marker=dict(size=8, symbol="circle")
                            ),
                            secondary_y=False
                        )
                        summary_data.append({
                            "槽名": loc_name,
                            "軸": "左軸",
                            "項目名": item_display_map[selected_item_y1],
                            "最新値": f"{valid_y1.iloc[-1]['Value_Num']:.2f}".rstrip('0').rstrip('.'),
                            "平均値": f"{valid_y1['Value_Num'].mean():.2f}".rstrip('0').rstrip('.'),
                            "最小値": f"{valid_y1['Value_Num'].min():.2f}".rstrip('0').rstrip('.'),
                            "最大値": f"{valid_y1['Value_Num'].max():.2f}".rstrip('0').rstrip('.')
                        })

                if selected_item_y2:
                    item_y2_row = df_item[(
                        df_item["Loc_ID"] == loc_id
                    ) & (
                        df_item["Sheet_Type"] == selected_sheet_type
                    ) & (
                        df_item["Item_Name"] == selected_item_y2
                    )]

                    if not item_y2_row.empty:
                        item_y2_id = item_y2_row.iloc[0]["Item_ID"]
                        rec_y2 = df_rec[df_rec["Item_ID"] == item_y2_id][["Report_ID", "Value_Num"]]
                        df_plot_y2 = site_reps_filtered.merge(rec_y2, on="Report_ID", how="left")

                        valid_y2 = df_plot_y2.dropna(subset=["Value_Num"])
                        if not valid_y2.empty:
                            has_data = True
                            fig.add_trace(
                                go.Scatter(
                                    x=df_plot_y2["Date"],
                                    y=df_plot_y2["Value_Num"],
                                    name=f"【{loc_name}】{item_display_map[selected_item_y2]} (右軸)",
                                    mode="lines+markers",
                                    connectgaps=False,
                                    line=dict(color=color_cfg["y2"], width=3),
                                    marker=dict(size=8, symbol="diamond")
                                ),
                                secondary_y=True
                            )
                            summary_data.append({
                                "槽名": loc_name,
                                "軸": "右軸",
                                "項目名": item_display_map[selected_item_y2],
                                "最新値": f"{valid_y2.iloc[-1]['Value_Num']:.2f}".rstrip('0').rstrip('.'),
                                "平均値": f"{valid_y2['Value_Num'].mean():.2f}".rstrip('0').rstrip('.'),
                                "最小値": f"{valid_y2['Value_Num'].min():.2f}".rstrip('0').rstrip('.'),
                                "最大値": f"{valid_y2['Value_Num'].max():.2f}".rstrip('0').rstrip('.')
                            })

            if not has_data:
                st.warning("選択した期間・槽・項目には数値データが存在しません。")
            else:
                title_locs_str = " vs ".join([graph_loc_options[lid] for lid in selected_graph_locs])
                title_text = f"【{site_options[selected_site_id]}】 ({title_locs_str}) 水質変化比較グラフ"

                fig.update_layout(
                    title_text=title_text,
                    xaxis_title="点検日 (Date)",
                    template="plotly_white",
                    hovermode="x unified",
                    legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
                )

                fig.update_yaxes(title_text=f"<b>{item_display_map[selected_item_y1]}</b>", secondary_y=False, title_font=dict(color="#2E7D32"))
                if selected_item_y2:
                    fig.update_yaxes(title_text=f"<b>{item_display_map[selected_item_y2]}</b>", secondary_y=True, title_font=dict(color="#1565C0"))

                st.plotly_chart(fig, use_container_width=True)
                st.caption("📸 グラフ右上のカメラアイコンをタップすると、グラフをPNG画像としてワンクリック保存できます。")

                if summary_data:
                    st.markdown("##### 📈 選択項目の水質統計サマリー (最新値・平均値・最小値・最大値)")
                    df_sum = pd.DataFrame(summary_data)
                    st.dataframe(df_sum, use_container_width=True)

# ------------------------------------------
# TAB 3: エクセル帳票一括ダウンロード (原本仕様2段ヘッダー出力 + 時系列AI水質診断)
# ------------------------------------------
with tab3:
    st.markdown("<div class='sub-header'>📑 全点検データの一括エクセル化 (原本仕様2段ヘッダー帳票出力)</div>", unsafe_allow_html=True)
    st.write("アップロードされた原本帳票と同じ「1行目: 槽名 / 2行目: 項目名」の2段ヘッダーおよび並び順で、一括エクセルファイルをダウンロードできます。")

    site_locs = df_loc[df_loc["Site_ID"] == selected_site_id].sort_values("Display_Order")
    loc_id_to_name = dict(zip(site_locs["Loc_ID"], site_locs["Loc_Name"]))
    items_ordered = df_item[(df_item["Loc_ID"].isin(site_locs["Loc_ID"])) & (df_item["Sheet_Type"] == selected_sheet_type)].sort_values("Display_Order")

    if items_ordered.empty:
        st.warning("該当する点検データが存在しません。")
    else:
        reports = df_rep[(df_rep["Site_ID"] == selected_site_id) & (df_rep["Sheet_Type"] == selected_sheet_type)].copy()
        reports["Date_dt"] = pd.to_datetime(reports["Date"], errors="coerce")
        reports = reports.sort_values("Date_dt").reset_index(drop=True)

        records = df_rec[df_rec["Report_ID"].isin(reports["Report_ID"])].copy()
        rec_dict = records.groupby(["Report_ID", "Item_ID"])["Value"].first().to_dict()

        col_defs = []
        for _, item_row in items_ordered.iterrows():
            l_name = loc_id_to_name.get(item_row["Loc_ID"], "")
            col_defs.append((item_row["Item_ID"], l_name, item_row["Item_Name"]))

        data_rows = []
        for _, rep_row in reports.iterrows():
            rep_id = rep_row["Report_ID"]
            r_date = rep_row["Date"]
            r_note = rep_row["Notes"] if pd.notna(rep_row["Notes"]) else ""
            row_vals = {"Date": r_date, "Notes": r_note}
            for item_id, l_name, i_name in col_defs:
                row_vals[item_id] = rec_dict.get((rep_id, item_id), "")
            data_rows.append(row_vals)

        df_raw = pd.DataFrame(data_rows)

        # 出力対象年の選択ドロップダウン (デフォルト: 全データ)
        available_years = sorted(reports["Date_dt"].dt.year.dropna().unique().astype(int).tolist(), reverse=True)
        year_options = ["全データ"] + [f"{y}年" for y in available_years]

        selected_year_opt = st.selectbox("📅 出力対象年を選択", options=year_options, index=0)

        if selected_year_opt == "全データ":
            df_export = df_raw.copy()
            file_year_str = "全データ"
        else:
            sel_y = int(selected_year_opt.replace("年", ""))
            df_export = df_raw[pd.to_datetime(df_raw["Date"], errors="coerce").dt.year == sel_y].copy()
            file_year_str = f"{sel_y}年"

        # 🤖 AI水質診断 ＆ 時系列傾向分析コメント (日付範囲指定トレンド検出)
        def get_ai_analysis_comment_advanced(selected_site_id, selected_sheet_type, df_export_data, col_defs_list):
            if df_export_data.empty:
                return ""
            
            latest_row = df_export_data.iloc[-1]
            latest_date = latest_row["Date"]
            latest_note = str(latest_row["Notes"]).strip() if pd.notna(latest_row["Notes"]) else ""
            salient_bullets = []
            
            for item_id, loc_name, item_name in col_defs_list:
                series_vals = df_export_data[item_id].dropna().astype(str)
                series_dates = df_export_data.loc[series_vals.index, "Date"].tolist()
                
                cleaned_nums = []
                valid_dates = []
                for idx_v, v in enumerate(series_vals):
                    v_clean = str(v).replace(",", "").replace("%", "").replace("<", "").replace(">", "").replace("未満", "").replace("以上", "").strip()
                    try:
                        f_v = float(v_clean)
                        cleaned_nums.append(f_v)
                        valid_dates.append(series_dates[idx_v])
                    except ValueError:
                        pass
                
                if len(cleaned_nums) >= 2:
                    mean_v = np.mean(cleaned_nums)
                    std_v = np.std(cleaned_nums)
                    l_v_num = cleaned_nums[-1]
                    latest_raw_v = str(df_export_data.iloc[-1].get(item_id, ""))
                    
                    # トレンド区間検出 (直近連続変化または最大傾き区間)
                    trend_msg = ""
                    tr_suffix = ""
                    if len(cleaned_nums) >= 3:
                        # 1. 連続増減チェック
                        inc_count = 0
                        dec_count = 0
                        for i in range(len(cleaned_nums)-1, 0, -1):
                            if cleaned_nums[i] > cleaned_nums[i-1]:
                                if dec_count > 0: break
                                inc_count += 1
                            elif cleaned_nums[i] < cleaned_nums[i-1]:
                                if inc_count > 0: break
                                dec_count += 1
                            else:
                                break
                        
                        unit_str = "%" if "割合" in item_name or "率" in item_name else (" mg/L" if "SS" in item_name or "BOD" in item_name or "COD" in item_name or "DO" in item_name else "")
                        
                        if inc_count >= 3:
                            s_d = valid_dates[-(inc_count+1)]
                            e_d = valid_dates[-1]
                            trend_msg = f"<b>{s_d}〜{e_d}</b> にかけて {inc_count+1} 回連続で上昇傾向（{cleaned_nums[-(inc_count+1)]:.1f} ➔ {l_v_num:.1f}{unit_str}）"
                            tr_suffix = f"（{s_d}〜{e_d} にかけて {inc_count+1} 回連続で上昇傾向）"
                        elif dec_count >= 3:
                            s_d = valid_dates[-(dec_count+1)]
                            e_d = valid_dates[-1]
                            trend_msg = f"<b>{s_d}〜{e_d}</b> にかけて {dec_count+1} 回連続で低下傾向（{cleaned_nums[-(dec_count+1)]:.1f} ➔ {l_v_num:.1f}{unit_str}）"
                            tr_suffix = f"（{s_d}〜{e_d} にかけて {dec_count+1} 回連続で低下傾向）"
                        elif len(cleaned_nums) >= 4:
                            # 2. 直近4回での全体の振れ幅・区間チェック
                            diff_pct = (l_v_num - cleaned_nums[-4]) / max(0.1, abs(cleaned_nums[-4]))
                            s_d = valid_dates[-4]
                            e_d = valid_dates[-1]
                            if diff_pct >= 0.25:
                                trend_msg = f"<b>{s_d}〜{e_d}</b> にかけて上昇傾向（{cleaned_nums[-4]:.1f} ➔ {l_v_num:.1f}{unit_str}）"
                                tr_suffix = f"（{s_d}〜{e_d} にかけて上昇傾向）"
                            elif diff_pct <= -0.25:
                                trend_msg = f"<b>{s_d}〜{e_d}</b> にかけて低下傾向（{cleaned_nums[-4]:.1f} ➔ {l_v_num:.1f}{unit_str}）"
                                tr_suffix = f"（{s_d}〜{e_d} にかけて低下傾向）"

                    # 1. SVI 判定
                    if item_name == "SVI":
                        if l_v_num > 150:
                            salient_bullets.append(f"<b>【{loc_name} SVI (汚泥膨化警報)】</b>: 最新値 <b>{l_v_num:.1f} mL/g</b>{tr_suffix}（全期間平均: {mean_v:.1f} mL/g）。指標値 150 mL/g を超えておりバルキング（膨化）傾向が見られます。DO管理および返送汚泥率の調整をご検討ください。")
                        elif l_v_num <= 150:
                            salient_bullets.append(f"<b>【{loc_name} SVI (汚泥沈降性)】</b>: 最新値 <b>{l_v_num:.1f} mL/g</b>{tr_suffix}（全期間平均: {mean_v:.1f} mL/g）。理想的な沈降範囲（150 mL/g以下）に収まっており、固液分離は極めて良好です。")
                    
                    # 2. BOD除去率 判定
                    elif item_name == "BOD除去率":
                        is_over = "以上" in str(latest_raw_v)
                        disp_str = f"{l_v_num:.1f}%以上" if is_over else f"{l_v_num:.1f}%"
                        if l_v_num >= 90:
                            salient_bullets.append(f"<b>【{loc_name} BOD除去率 (処理性能)】</b>: 最新値 <b>{disp_str}</b>{tr_suffix}（全期間平均: {mean_v:.1f}%）。高い除去効果を維持しており、微生物活性・処理プロセスは非常に安定しています。")
                        else:
                            salient_bullets.append(f"<b>【{loc_name} BOD除去率 (処理低下注意)】</b>: 最新値 <b>{disp_str}</b>{tr_suffix}（全期間平均: {mean_v:.1f}%）。除去率低下傾向がうかがえます。原水負荷または曝気量の調整をおすすめします。")
                    
                    # 3. 処理水/放流SS 判定
                    elif "処理" in loc_name and "SS" in item_name:
                        if l_v_num > (mean_v + 2.0 * std_v) and mean_v > 0:
                            salient_bullets.append(f"<b>【{loc_name} {item_name} (流出注意)】</b>: 最新値 <b>{latest_raw_v} mg/L</b>{tr_suffix}（全期間平均: {mean_v:.1f} mg/L）。全期間平均を上回って推移しています。沈殿槽フロックの流出有無をご確認ください。")
                        elif l_v_num == 0 or "未満" in str(latest_raw_v):
                            salient_bullets.append(f"<b>【{loc_name} {item_name} (透視度良好)】</b>: 最新値 <b>{latest_raw_v} mg/L</b>（全期間平均: {mean_v:.1f} mg/L）。透視度の高い清澄な処理水が得られています。")
                    
                    # 4. pH 異常判定 (6.0〜8.5 外れ)
                    elif item_name in ["pH", "ph"]:
                        if l_v_num < 6.0 or l_v_num > 8.5:
                            salient_bullets.append(f"<b>【{loc_name} pH (中性域外れ)】</b>: 最新値 <b>{l_v_num:.2f}</b>{tr_suffix}（全期間平均: {mean_v:.2f}）。管理標準域 (6.0〜8.5) を外れています。原水流入およびアルカリ/酸注入状態をご確認ください。")
                    
                    # 5. 期間指定トレンド検出項目
                    elif trend_msg:
                        unit_str = "%" if "割合" in item_name or "率" in item_name else (" mg/L" if "SS" in item_name or "BOD" in item_name or "COD" in item_name or "DO" in item_name else "")
                        salient_bullets.append(f"<b>【{loc_name} {item_name} (期間推移)】</b>: {trend_msg}（全期間平均: {mean_v:.2f}{unit_str}）。")
                    
                    # 6. その他の大幅乖離項目 (2.5σ 超え)
                    elif std_v > 0 and abs(l_v_num - mean_v) > 2.5 * std_v:
                        unit_str = "%" if "割合" in item_name or "率" in item_name else ""
                        salient_bullets.append(f"<b>【{loc_name} {item_name} (変動検出)】</b>: 最新値 <b>{latest_raw_v}</b>（全期間平均: {mean_v:.2f}{unit_str}）。通常推移範囲からの変動が検出されました。")

            if not salient_bullets:
                salient_bullets.append("<b>【水質全般 安定維持】</b>: 全測定項目が過去の通常変動範囲内で推移しており、処理プロセスは極めて良好に維持されています。")
                
            note_html = f" <span style='color:#555555; font-size:13px;'>[現場メモ: 「{latest_note}」]</span>" if latest_note and latest_note != "nan" else ""
            bullets_html = "".join([f"<li style='margin-bottom:6px;'>{b}</li>" for b in salient_bullets[:5]])
            
            html_content = f"""
<div style='background-color:#F1F8E9; border-left:5px solid #66BB6A; border-radius:8px; padding:14px; margin-top:12px; margin-bottom:15px;'>
    <div style='font-size:15px; font-weight:bold; color:#2E7D32; margin-bottom:8px;'>
        🤖 AI水質診断 ＆ 傾向分析コメント（最新点検日: {latest_date}）{note_html}
    </div>
    <ul style='margin:0; padding-left:20px; font-size:14px; color:#333333; line-height:1.7;'>
        {bullets_html}
    </ul>
</div>
"""
            return html_content

        ai_comment_html = get_ai_analysis_comment_advanced(selected_site_id, selected_sheet_type, df_export, col_defs)
        st.markdown(ai_comment_html, unsafe_allow_html=True)

        mi_cols = [("基本情報", "日付")] + [(loc_name, item_name) for _, loc_name, item_name in col_defs] + [("基本情報", "備考")]
        raw_col_keys = ["Date"] + [item_id for item_id, _, _ in col_defs] + ["Notes"]

        df_preview = df_export[raw_col_keys].copy()
        df_preview.columns = pd.MultiIndex.from_tuples(mi_cols)

        st.markdown(f"##### プレビュー ({selected_year_opt} - 2段ヘッダー構造)")
        st.dataframe(df_preview.tail(10) if len(df_preview) > 10 else df_preview, use_container_width=True)

        def generate_formatted_excel(selected_site_id, selected_sheet_type, site_name, df_data_source):
            output = io.BytesIO()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"{selected_sheet_type}_集計"
            ws.append([f"◆ {site_name} 【{selected_sheet_type}】 {file_year_str} 点検データ一覧表"])
            ws.append([f"出力日時: {datetime.datetime.now().strftime('%Y/%m/%d %H:%M')}"])
            ws.append([])

            header1 = ["日付"] + [loc_name for _, loc_name, _ in col_defs] + ["備考"]
            header2 = ["日付"] + [item_name for _, _, item_name in col_defs] + ["備考"]
            ws.append(header1)
            ws.append(header2)

            for _, r in df_data_source.iterrows():
                row_arr = [r["Date"]] + [r[item_id] for item_id, _, _ in col_defs] + [r["Notes"]]
                ws.append(row_arr)

            total_cols = len(header1)
            ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
            ws.merge_cells(start_row=4, start_column=total_cols, end_row=5, end_column=total_cols)

            curr_start = 2
            for c_idx in range(2, total_cols):
                loc_curr = header1[c_idx - 1]
                loc_next = header1[c_idx] if c_idx < total_cols - 1 else None
                if loc_curr != loc_next:
                    if curr_start < c_idx:
                        ws.merge_cells(start_row=4, start_column=curr_start, end_row=4, end_column=c_idx)
                    curr_start = c_idx + 1

            fill_header1 = PatternFill(start_color="43A047", end_color="43A047", fill_type="solid")
            fill_header2 = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
            font_header1 = Font(name="メイリオ", size=11, bold=True, color="FFFFFF")
            font_header2 = Font(name="メイリオ", size=10, bold=True, color="2E7D32")
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style="thin", color="A5D6A7"),
                right=Side(style="thin", color="A5D6A7"),
                top=Side(style="thin", color="A5D6A7"),
                bottom=Side(style="thin", color="A5D6A7")
            )

            ws.cell(row=1, column=1).font = Font(name="メイリオ", size=14, bold=True, color="2E7D32")
            ws.cell(row=2, column=1).font = Font(name="メイリオ", size=10, italic=True, color="666666")

            for col in range(1, total_cols + 1):
                cell1 = ws.cell(row=4, column=col)
                cell2 = ws.cell(row=5, column=col)
                cell1.fill = fill_header1
                cell1.font = font_header1
                cell1.alignment = align_center
                cell1.border = thin_border
                cell2.fill = fill_header2
                cell2.font = font_header2
                cell2.alignment = align_center
                cell2.border = thin_border

            for row_idx in range(6, ws.max_row + 1):
                for col_idx in range(1, total_cols + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(name="メイリオ", size=10)
                    cell.border = thin_border
                    if col_idx == 1:
                        cell.alignment = Alignment(horizontal="center")
                    elif col_idx == total_cols:
                        cell.alignment = Alignment(horizontal="left")
                    else:
                        cell.alignment = Alignment(horizontal="right")

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row >= 4 and cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

            wb.save(output)
            output.seek(0)
            return output

        excel_data = generate_formatted_excel(selected_site_id, selected_sheet_type, site_options[selected_site_id], df_export)
        st.download_button(
            label=f"📥 【{site_options[selected_site_id]}】 ({selected_sheet_type} - {file_year_str}) エクセル帳票をダウンロード",
            data=excel_data,
            file_name=f"{site_options[selected_site_id]}_{selected_sheet_type}_{file_year_str}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

# ------------------------------------------
# TAB 4: データベース全件閲覧 (要件⑦)
# ------------------------------------------
with tab4:
    st.markdown("<div class='sub-header'>🗃️ マスターデータ ＆ 全蓄積レコード閲覧</div>", unsafe_allow_html=True)
    m_tab1, m_tab2, m_tab3, m_tab4 = st.tabs(["工場マスター", "測定箇所マスター", "項目マスター", "全蓄積レコード"])
    with m_tab1:
        st.dataframe(df_site, use_container_width=True)
    with m_tab2:
        st.dataframe(df_loc[df_loc["Site_ID"] == selected_site_id], use_container_width=True)
    with m_tab3:
        st.dataframe(df_item[(df_item["Loc_ID"].isin(df_loc[df_loc["Site_ID"] == selected_site_id]["Loc_ID"])) & (df_item["Sheet_Type"] == selected_sheet_type)], use_container_width=True)
    with m_tab4:
        cur_site_reps = df_rep[(df_rep["Site_ID"] == selected_site_id) & (df_rep["Sheet_Type"] == selected_sheet_type)]
        cur_site_recs = df_rec[df_rec["Report_ID"].isin(cur_site_reps["Report_ID"])]
        st.dataframe(cur_site_recs.tail(200), use_container_width=True)
